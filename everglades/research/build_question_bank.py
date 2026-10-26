#!/usr/bin/env python3
"""Builds Everglades_Question_Bank.xlsx from questions.py and the site taxonomy."""
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "build"))
from questions import BUYERS, BY_CATEGORY, CATEGORY_BUYER, CORE, TOPICS  # noqa: E402
from taxonomy import CATEGORIES, SECTIONS  # noqa: E402

OUT = os.path.join(HERE, "Everglades_Question_Bank.xlsx")
CAT = {c["n"]: c for c in CATEGORIES}
SEC = {k: n for k, n, _ in SECTIONS}
ALL_TOPICS = TOPICS + ["Services Offered"]

GREEN = "0B6B1F"
DEEP = "06381A"
YELLOW_FILL = PatternFill("solid", fgColor="FFF3C4")
HEAD_FILL = PatternFill("solid", fgColor=DEEP)
BAND_FILL = PatternFill("solid", fgColor="F6F3EA")
F = "Arial"
thin = Side(style="thin", color="DAD3C0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def topic_for(q, sec):
    s = q.lower()
    if " vs. " in s or "difference" in s or "better" in s:
        return "Comparisons"
    if "cost" in s or "how much" in s or "grants" in s:
        return "Cost"
    if "permit" in s or "required" in s or "have to be" in s or "responsible" in s or "qualification" in s or "training" in s:
        return "Permits & Rules"
    if s.startswith(("how long", "how fast", "how often", "how soon", "when")):
        return "Timing"
    if "fema" in s or "storm" in s or "hurricane" in s or "disaster" in s:
        return "Storm & FEMA"
    if "army corps" in s or "military" in s or "municipal" in s or "solicitation" in s:
        return "Government Buying"
    if "out of state" in s or "nationwide" in s or "remotely" in s:
        return "Projects Anywhere"
    if s.startswith(("do you", "can you", "can i rent")):
        return "Services Offered"
    if sec in ("water", "cleanup") or "wetland" in s or "invasive" in s or "habitat" in s or "native" in s:
        return "Environmental"
    return "Equipment & Methods"


def needs_matt(q, topic):
    s = q.lower()
    if topic in ("Cost", "Timing", "Working With Us", "Government Buying", "Projects Anywhere", "Services Offered"):
        return "Yes"
    if s.startswith(("do you", "can you", "are you", "how do you", "how big")) or " your " in s:
        return "Yes"
    return "No"


def video_ok(q):
    s = q.lower()
    return "Yes" if s.startswith(("what is", "what are", "how do", "how is", "how long", "how wide", "how deep", "can ", "what happens")) \
        or " vs. " in s or "difference" in s or "or " in s else "No"


def build_rows():
    rows, seen = [], set()

    def add(q, topic, buyer, tier, cat=None):
        key = re.sub(r"[^a-z0-9 ]", "", q.lower()).strip()
        if key in seen:
            return
        seen.add(key)
        c = CAT.get(cat)
        rows.append({
            "q": q, "topic": topic, "buyer": buyer, "tier": tier,
            "cat_n": c["n"] if c else "", "cat": c["name"] if c else "All capabilities",
            "sec": SEC[c["sec"]] if c else "Cross-cutting",
            "matt": needs_matt(q, topic), "video": video_ok(q),
            "page": f"/capabilities/{c['slug']}/" if c else "/questions/",
        })

    for q, topic, buyer in CORE:
        add(q, topic, buyer, 1)
    for sk, _, _ in SECTIONS:
        for c in [c for c in CATEGORIES if c["sec"] == sk]:
            for i, q in enumerate(BY_CATEGORY[c["n"]]):
                add(q, topic_for(q, sk), CATEGORY_BUYER[c["n"]], 2 if i < 4 else 3, c["n"])
    for i, r in enumerate(rows, 1):
        r["id"] = f"Q{i:03d}"
    return rows


def style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 30


def main():
    rows = build_rows()
    wb = Workbook()

    # ------------------------------------------------------------ Start Here
    ws = wb.active
    ws.title = "Start Here"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 95
    lines = [
        ("title", "Everglades Restorations & Land Clearing — Question Bank"),
        ("sub", "The questions buyers ask Google, YouTube, voice assistants and AI tools. Each answered question becomes its own page, video Short or FAQ block."),
        ("gap", ""),
        ("h", "How to use it"),
        ("kv", ("1. Meeting Top 50", "Start here at the meeting. These are the Tier 1 questions only Matt and Sam can answer with real numbers and experience.")),
        ("kv", ("2. Question Bank", "Every question, sorted by topic, capability, buyer and tier. Filter any column.")),
        ("kv", ("3. Fill in", "Type answers in the yellow \"Matt's Answer\" column, in plain talk, the way you'd say it on the phone. We turn them into finished pages.")),
        ("kv", ("4. Summary", "Counts by tier, topic, section and buyer. Updates automatically as Status changes.")),
        ("gap", ""),
        ("h", "Tiers"),
        ("kv", ("Tier 1 — Core", "Cross-cutting questions (cost, permits, timing, storm/FEMA, government buying). Build first, answer deeply.")),
        ("kv", ("Tier 2 — Category FAQ", "Top questions for each of the 38 capability categories. Shown on each category page.")),
        ("kv", ("Tier 3 — Service FAQ", "Deeper questions. Built only as each service page gets real content and jobs.")),
        ("kv", ("Tier 4 — Job-born", "Not listed here: questions answered on camera in each new job video. Grows every week.")),
        ("gap", ""),
        ("h", "Columns"),
        ("kv", ("Needs Matt's Answer", "Yes = needs real numbers, experience or company facts. No = a definition we can draft for review.")),
        ("kv", ("Video Short?", "Yes = a good 30–60 second spoken answer on YouTube (shows up in AI answers too).")),
        ("kv", ("Status", "Open → Answered → Drafted → Live. Change it from the dropdown.")),
        ("kv", ("Page", "Where the answer will live on the site.")),
        ("gap", ""),
        ("h", "Example of a finished answer"),
        ("kv", ("Question", "How much does land clearing cost per acre?")),
        ("kv", ("Matt's Answer (plain talk)", "Light brush, we can mulch it for about $X,XXX an acre. Thick stuff with big trees and haul-off runs $X,XXX to $X,XXX. Wet ground or permits add to it. We'll come look and give you a firm number.")),
        ("kv", ("Becomes the short answer (voice)", "Land clearing usually runs $X,XXX to $X,XXX per acre. Light brush mulched in place is on the low end; heavy trees, haul-off, wet ground and permits push it higher. A site visit gives a firm price.")),
        ("gap", ""),
        ("h", "Shape of every answer page"),
        ("kv", ("1. Short answer", "40–60 words, about 20 seconds read aloud. This is what voice assistants and AI answers lift.")),
        ("kv", ("2. Detail", "150–400 words: real numbers, steps, rules.")),
        ("kv", ("3. Proof", "A Short video answer, plus links to the related service and real job pages.")),
        ("gap", ""),
        ("note", f"{len(rows)} questions · {sum(1 for r in rows if r['tier'] == 1)} Tier 1 · {sum(1 for r in rows if r['tier'] == 2)} Tier 2 · {sum(1 for r in rows if r['tier'] == 3)} Tier 3. Room to grow to 3,000+ as service pages and jobs are added; quality first, per Google's scaled-content rules."),
    ]
    r = 2
    for kind, val in lines:
        if kind == "title":
            ws.cell(row=r, column=2, value=val).font = Font(name=F, bold=True, size=18, color=DEEP)
        elif kind == "sub":
            c = ws.cell(row=r, column=2, value=val)
            c.font = Font(name=F, size=11, color="4A4A45")
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        elif kind == "h":
            ws.cell(row=r, column=2, value=val).font = Font(name=F, bold=True, size=12, color=GREEN)
        elif kind == "kv":
            a = ws.cell(row=r, column=2, value=val[0])
            a.font = Font(name=F, bold=True, size=10)
            a.alignment = Alignment(vertical="top", wrap_text=True)
            b = ws.cell(row=r, column=3, value=val[1])
            b.font = Font(name=F, size=10)
            b.alignment = Alignment(vertical="top", wrap_text=True)
            if val[0].startswith("Matt's Answer"):
                b.fill = YELLOW_FILL
            ws.row_dimensions[r].height = 30 if len(val[1]) > 95 else 16
        elif kind == "note":
            c = ws.cell(row=r, column=2, value=val)
            c.font = Font(name=F, italic=True, size=10, color="4A4A45")
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        r += 1

    # ------------------------------------------------------------ Question Bank
    qb = wb.create_sheet("Question Bank")
    headers = ["ID", "Question", "Topic", "Cap #", "Capability", "Section", "Buyer", "Tier",
               "Needs Matt's Answer", "Video Short?", "Matt's Answer (plain talk)", "Status", "Page", "Notes"]
    widths = [7, 52, 18, 7, 34, 24, 22, 6, 11, 9, 50, 11, 38, 24]
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        qb.cell(row=1, column=i, value=h)
        qb.column_dimensions[get_column_letter(i)].width = w
    style_header(qb, 1, len(headers))
    for i, row in enumerate(rows, 2):
        vals = [row["id"], row["q"], row["topic"], row["cat_n"], row["cat"], row["sec"], row["buyer"], row["tier"],
                row["matt"], row["video"], None, "Open", row["page"], None]
        for j, v in enumerate(vals, 1):
            c = qb.cell(row=i, column=j, value=v)
            c.font = Font(name=F, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=j in (2, 5, 11, 14))
            c.border = BORDER
            if j == 11:
                c.fill = YELLOW_FILL
            elif i % 2 == 0:
                c.fill = BAND_FILL
    last = len(rows) + 1
    qb.freeze_panes = "C2"
    qb.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    dv = DataValidation(type="list", formula1='"Open,Answered,Drafted,Live,Skip"', allow_blank=False)
    qb.add_data_validation(dv)
    dv.add(f"L2:L{last}")

    # ------------------------------------------------------------ Meeting Top 50
    mt = wb.create_sheet("Meeting Top 50", 0)
    order = {t: i for i, t in enumerate(["Cost", "Timing", "Permits & Rules", "Storm & FEMA", "Government Buying",
                                         "Working With Us", "Projects Anywhere", "Equipment & Methods", "Environmental", "Comparisons"])}
    pick = [r for r in rows if r["tier"] == 1 and r["matt"] == "Yes"]
    pick.sort(key=lambda r: (order.get(r["topic"], 99), r["id"]))
    pick = pick[:50]
    mh = ["#", "ID", "Question", "Topic", "Buyer", "Matt's Answer (plain talk)"]
    mw = [5, 7, 60, 18, 22, 70]
    mt.cell(row=1, column=1, value="Meeting Top 50 — questions only Matt & Sam can answer").font = Font(name=F, bold=True, size=14, color=DEEP)
    mt.cell(row=2, column=1, value="Answer in plain talk, like on the phone. Yellow column = type here. The same IDs are in the Question Bank tab.").font = Font(name=F, italic=True, size=10, color="4A4A45")
    for i, (h, w) in enumerate(zip(mh, mw), 1):
        mt.cell(row=4, column=i, value=h)
        mt.column_dimensions[get_column_letter(i)].width = w
    style_header(mt, 4, len(mh))
    for i, row in enumerate(pick, 5):
        for j, v in enumerate([i - 4, row["id"], row["q"], row["topic"], row["buyer"], None], 1):
            c = mt.cell(row=i, column=j, value=v)
            c.font = Font(name=F, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = BORDER
            if j == 6:
                c.fill = YELLOW_FILL
        mt.row_dimensions[i].height = 42
    mt.freeze_panes = "C5"

    # ------------------------------------------------------------ Summary (formulas)
    sm = wb.create_sheet("Summary")
    sm.sheet_view.showGridLines = False
    rng = lambda col: f"'Question Bank'!${col}$2:${col}${last}"  # noqa: E731
    sm.column_dimensions["A"].width = 30
    for col in "BCDEFG":
        sm.column_dimensions[col].width = 13
    sm["A1"] = "Question Bank Summary"
    sm["A1"].font = Font(name=F, bold=True, size=14, color=DEEP)
    sm["A2"] = "All counts are live formulas over the Question Bank tab."
    sm["A2"].font = Font(name=F, italic=True, size=10, color="4A4A45")

    def table(start, title, labels, label_col, extra_cols):
        sm.cell(row=start, column=1, value=title)
        hdr = [title.split(" by ")[-1]] + [h for h, _ in extra_cols]
        for j, h in enumerate(hdr, 1):
            sm.cell(row=start, column=j, value=h)
        style_header(sm, start, len(hdr))
        r0 = start + 1
        for i, lab in enumerate(labels):
            rr = r0 + i
            sm.cell(row=rr, column=1, value=lab).font = Font(name=F, size=10)
            for j, (_, fn) in enumerate(extra_cols, 2):
                c = sm.cell(row=rr, column=j, value=fn(rng(label_col), f"$A{rr}"))
                c.font = Font(name=F, size=10)
                c.alignment = Alignment(horizontal="center")
        tr = r0 + len(labels)
        sm.cell(row=tr, column=1, value="Total").font = Font(name=F, bold=True, size=10)
        for j in range(2, len(extra_cols) + 2):
            colL = get_column_letter(j)
            c = sm.cell(row=tr, column=j, value=f"=SUM({colL}{r0}:{colL}{tr - 1})")
            c.font = Font(name=F, bold=True, size=10)
            c.alignment = Alignment(horizontal="center")
        return tr + 2

    cols = [
        ("Questions", lambda R, a: f"=COUNTIF({R},{a})"),
        ("Tier 1", lambda R, a: f"=COUNTIFS({R},{a},{rng('H')},1)"),
        ("Tier 2", lambda R, a: f"=COUNTIFS({R},{a},{rng('H')},2)"),
        ("Tier 3", lambda R, a: f"=COUNTIFS({R},{a},{rng('H')},3)"),
        ("Needs Matt", lambda R, a: f"=COUNTIFS({R},{a},{rng('I')},\"Yes\")"),
        ("Answered+", lambda R, a: f"=COUNTIFS({R},{a},{rng('L')},\"<>Open\")"),
    ]
    nxt = table(4, "Questions by Topic", ALL_TOPICS, "C", cols)
    nxt = table(nxt, "Questions by Section", ["Cross-cutting"] + [n for _, n, _ in SECTIONS], "F", cols)
    table(nxt, "Questions by Buyer", BUYERS, "G", cols)

    # ------------------------------------------------------------ Sources
    so = wb.create_sheet("Sources")
    so.column_dimensions["A"].width = 34
    so.column_dimensions["B"].width = 90
    so["A1"] = "Where questions come from (and where to find more)"
    so["A1"].font = Font(name=F, bold=True, size=14, color=DEEP)
    src = [
        ("Matt & Sam's phone calls", "Best source. The questions customers actually ask. Capture at the meeting and keep a running list."),
        ("Google 'People also ask' + autocomplete", "Type each service into Google and record the questions that appear."),
        ("Reddit & YouTube comments", "Real wording from homeowners and contractors on land clearing, ponds, permits and FEMA debris."),
        ("Government bid Q&A addenda", "Agencies publish answers to bidder questions; they show what procurement officers ask."),
        ("Search Console (after launch)", "Shows the real queries bringing people to the site. Add new questions monthly."),
        ("Google: FAQPage structured data", "https://developers.google.com/search/docs/appearance/structured-data/faqpage — FAQ rich results are limited to well-known government and health sites; markup still helps machines read Q&A."),
        ("Search Engine Land: rise and fall of FAQ schema", "https://searchengineland.com/faq-schema-rise-fall-seo-today-463993"),
        ("Google spam policies: scaled content abuse", "https://developers.google.com/search/docs/essentials/spam-policies — mass-produced, low-value pages are penalized; build answers with real expertise."),
        ("AI citation research (vendor studies)", "https://docs.peec.ai/research/most-cited-domains-by-engine — YouTube and Reddit are cited across AI engines; findings shift quickly."),
    ]
    for i, (a, b) in enumerate(src, 3):
        so.cell(row=i, column=1, value=a).font = Font(name=F, bold=True, size=10)
        c = so.cell(row=i, column=2, value=b)
        c.font = Font(name=F, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        so.row_dimensions[i].height = 30

    wb.move_sheet("Start Here", offset=-wb.index(wb["Start Here"]))
    wb.active = 0
    wb.save(OUT)
    print(f"{len(rows)} questions → {OUT}")


if __name__ == "__main__":
    main()
